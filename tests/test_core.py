"""Проверка без железа: API-клиент, синхронизация, SQLite, экспорт, графики, CLI - против имитатора ESP32.
Запуск:  python -m unittest discover -s tests -v
"""
import json
import os
import sqlite3
import sys
import tempfile
import time
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "tools"))

import matplotlib
matplotlib.use("Agg")

from barograph import cli, exporters, plotting, service
from barograph.api_client import BarographClient, DeviceError, parse_csv_text
from barograph.paths import Config
from barograph.storage import Storage
from barograph.timeutil import parse_datetime, parse_duration, tz_from_offset
from mock_esp32 import start_mock

TZ = tz_from_offset(10800)


class CoreTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.srv, cls.port, cls.dev = start_mock(0, hours=72)
        cls.tmp = tempfile.TemporaryDirectory()
        cls.client = BarographClient("127.0.0.1", cls.port, timeout=5)

    @classmethod
    def tearDownClass(cls):
        cls.srv.shutdown()
        cls.tmp.cleanup()

    def path(self, name):
        return os.path.join(self.tmp.name, name)

    # ---------------- API
    def test_01_api_basic(self):
        st = self.client.status()
        self.assertEqual(st["sensor"], "BME280")
        cur = self.client.current()
        self.assertEqual(cur["state"], "ok")
        self.assertTrue(self.client.forecast()["valid"])
        h = self.client.history(24, 200)
        self.assertLessEqual(h["count"], 200)
        m = self.client.measure(store=False)
        self.assertTrue(m["success"])
        self.assertEqual(self.client.get_settings()["record_interval_s"], 60)

    def test_02_settings_validation(self):
        with self.assertRaises(DeviceError) as cm:
            self.client.set_settings(measure_interval_s=1, record_interval_s=60)
        self.assertIn("invalid_interval", str(cm.exception))
        r = self.client.set_settings(measure_interval_s=30, record_interval_s=300)
        self.assertEqual(r["measure_interval_s"], 30)
        self.client.set_settings(measure_interval_s=10, record_interval_s=60)

    def test_03_wifi_scan_and_time(self):
        nets = self.client.wifi_scan(wait_s=10, poll_s=0.3)
        self.assertEqual(len(nets), 3)
        self.assertTrue(self.client.set_time()["success"])

    def test_04_no_connection(self):
        c = BarographClient("127.0.0.1", 1, timeout=1)
        with self.assertRaises(DeviceError):
            c.status()

    def test_05_fallback_host(self):
        c = BarographClient("no-such-host.invalid", self.port, timeout=2, fallback_hosts=["127.0.0.1"])
        self.assertEqual(c.status()["fw_version"], "1.0.0")
        self.assertEqual(c.active_host, "127.0.0.1")

    # ---------------- CSV
    def test_10_parse_csv(self):
        text = ("timestamp,date,time,pressure_mmHg,temperature_c,humidity\n"
                "1791097200,2026-10-04,12:00:00,755.42,18.7,61.4\n"
                "1791097260,2026-10-04,12:01:00,755.38,-3.2,\n"
                "garbage line\n1791097320,2026-10-04,12:0\n")
        rec = parse_csv_text(text)
        self.assertEqual(rec, [(1791097200, 755.42, 18.7, 61.4), (1791097260, 755.38, -3.2, None)])

    # ---------------- sync + storage
    def test_20_sync_and_dedup(self):
        db = Storage(self.path("t.db"))
        r1 = service.sync_from_device(self.client, db)
        self.assertEqual(r1.downloaded, 4320)
        self.assertEqual(r1.new, 4320)
        r2 = service.sync_from_device(self.client, db)          # повторная синхронизация - без дублей
        self.assertEqual(r2.new, 0)
        self.assertEqual(db.count(), 4320)
        # новые записи на устройстве
        self.dev.records.append((self.dev.records[-1][0] + 60, 756.0, 19.0, 50.0))
        r3 = service.sync_from_device(self.client, db)
        self.assertEqual(r3.new, 1)

    def test_21_query_resample_stats(self):
        db = Storage(self.path("t2.db"))
        service.sync_from_device(self.client, db)
        lo, hi = db.bounds()
        raw = db.query(lo, hi)
        hourly = db.query(lo, hi, bucket_s=3600)
        self.assertGreater(len(raw), len(hourly))
        self.assertLessEqual(len(hourly), 74)
        auto = db.query(lo, hi, max_points=500)
        self.assertLessEqual(len(auto), 520)
        st = db.stats(lo, hi)
        self.assertEqual(st["count"], len(raw))
        self.assertLess(st["pressure"]["min"], st["pressure"]["max"])

    def test_22_bmp_no_humidity(self):
        srv, port, dev = start_mock(0, hours=2, has_humidity=False)
        try:
            c = BarographClient("127.0.0.1", port)
            db = Storage(self.path("bmp.db"))
            service.sync_from_device(c, db)
            rows = db.query()
            self.assertTrue(all(r[3] is None for r in rows))
            st = db.stats()
            self.assertIsNone(st["humidity"]["min"])
            # экспорт и график без влажности не падают
            service.export_data(db, "xlsx", self.path("bmp.xlsx"), 0, 4_000_000_000, TZ)
            service.export_data(db, "png", self.path("bmp.png"), 0, 4_000_000_000, TZ)
        finally:
            srv.shutdown()

    # ---------------- экспорт
    def test_30_exports(self):
        db = Storage(self.path("e.db"))
        service.sync_from_device(self.client, db)
        lo, hi = db.bounds()
        n = service.export_data(db, "csv", self.path("e.csv"), lo, hi, TZ)
        self.assertEqual(n, db.count())
        # CSV можно загрузить обратно (формат совместим с файлом прошивки)
        db2 = Storage(self.path("e2.db"))
        parsed, new = service.import_csv_file(db2, self.path("e.csv"))
        self.assertEqual(new, db.count())
        self.assertEqual(db2.query()[0][0], db.query()[0][0])

        service.export_data(db, "csv", self.path("ru.csv"), lo, hi, TZ, excel_ru=True)
        with open(self.path("ru.csv"), encoding="utf-8-sig") as f:
            line = f.readlines()[1]
        self.assertIn(";", line)
        self.assertRegex(line, r"\d,\d")

        service.export_data(db, "json", self.path("e.json"), lo, hi, TZ, unit="hPa")
        with open(self.path("e.json"), encoding="utf-8") as f:
            doc = json.load(f)
        self.assertEqual(doc["pressure_unit"], "hPa")
        self.assertGreater(doc["data"][0]["pressure"], 900)
        self.assertEqual(doc["timezone_offset_s"], 10800)

        service.export_data(db, "sqlite", self.path("out.sqlite"), lo, hi, TZ, resample_s=3600)
        conn = sqlite3.connect(self.path("out.sqlite"))
        cnt = conn.execute("SELECT COUNT(*) FROM readings").fetchone()[0]
        conn.close()
        self.assertLessEqual(cnt, 74)

        service.export_data(db, "xlsx", self.path("e.xlsx"), lo, hi, TZ, resample_s=900)
        import openpyxl
        wb = openpyxl.load_workbook(self.path("e.xlsx"))
        self.assertEqual(wb.sheetnames, ["Данные", "Статистика", "Графики"])
        ws = wb["Данные"]
        self.assertGreater(ws.max_row, 100)
        self.assertEqual(ws["A2"].number_format, "DD.MM.YYYY HH:MM:SS")
        self.assertEqual(len(wb["Графики"]._charts), 3)

        service.export_data(db, "png", self.path("e.png"), lo, hi, TZ)
        self.assertGreater(os.path.getsize(self.path("e.png")), 10_000)

    def test_31_empty_period(self):
        db = Storage(self.path("empty.db"))
        for fmt in ("csv", "xlsx", "json", "sqlite", "png"):
            with self.assertRaises(exporters.ExportError):
                service.export_data(db, fmt, self.path(f"empty.{fmt}"), 0, 100, TZ)

    # ---------------- график с разрывом
    def test_40_plot_gap(self):
        from matplotlib.figure import Figure
        base = 1_790_000_000
        rows = [(base + i * 60, 750 + i * 0.01, 20.0, 50.0) for i in range(60)]
        rows += [(base + 20000 + i * 60, 752 + i * 0.01, 21.0, 52.0) for i in range(60)]
        fig = Figure()
        plotting.draw_figure(fig, rows, ["pressure", "temperature", "humidity"], TZ, "mmHg")
        self.assertEqual(len(fig.axes), 3)
        ys = fig.axes[0].lines[0].get_ydata()
        self.assertTrue(any(y != y for y in ys))               # NaN вставлен на месте разрыва

    # ---------------- вспомогательное
    def test_50_time_parsing(self):
        a = parse_datetime("04.10.2026 12:30", TZ)
        b = parse_datetime("2026-10-04 12:30", TZ)
        self.assertEqual(a, b)
        self.assertEqual(parse_datetime("04.10.2026", TZ, end=True) - parse_datetime("04.10.2026", TZ), 86399)
        with self.assertRaises(ValueError):
            parse_datetime("вчера", TZ)
        self.assertEqual(parse_duration("1h"), 3600)
        self.assertEqual(parse_duration("5m"), 300)
        self.assertEqual(parse_duration("0"), 0)

    def test_60_cli(self):
        db = self.path("cli.db")
        out = self.path("cli.csv")
        base = ["--host", "127.0.0.1", "--port", str(self.port), "--db", db]
        self.assertEqual(cli.main(base + ["sync"]), 0)
        self.assertEqual(cli.main(base + ["stats", "--hours", "24"]), 0)
        self.assertEqual(cli.main(base + ["export", "--format", "csv", "--out", out, "--hours", "6",
                                          "--resample", "5m", "--unit", "hPa"]), 0)
        self.assertTrue(os.path.exists(out))
        self.assertEqual(cli.main(base + ["status"]), 0)
        self.assertEqual(cli.main(["--host", "127.0.0.1", "--port", "1", "--db", db, "current"]), 2)


if __name__ == "__main__":
    unittest.main()
