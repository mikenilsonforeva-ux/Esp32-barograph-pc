"""Экспорт данных: CSV, Excel (xlsx), JSON. (SQLite - в Storage.export_sqlite, PNG - в plotting.)"""
import csv
import json

from .timeutil import fmt_ts, to_dt
from .units import convert_pressure, pressure_label

MAX_EXCEL_ROWS = 1_000_000          # лимит Excel - 1 048 576 строк


class ExportError(Exception):
    pass


def _check_rows(rows):
    if not rows:
        raise ExportError("Нет данных за выбранный период")


def _num(v, digits, decimal_comma=False):
    if v is None:
        return ""
    s = f"{v:.{digits}f}"
    return s.replace(".", ",") if decimal_comma else s


def export_csv(rows, path, tz, unit="mmHg", excel_ru=False) -> int:
    """CSV в формате файла прошивки: timestamp,date,time,pressure_<unit>,temperature_c,humidity.
    excel_ru=True: разделитель «;», десятичная запятая, UTF-8 с BOM - для русской локали Excel."""
    _check_rows(rows)
    delim = ";" if excel_ru else ","
    with open(path, "w", newline="", encoding="utf-8-sig" if excel_ru else "utf-8") as f:
        w = csv.writer(f, delimiter=delim)
        w.writerow(["timestamp", "date", "time", f"pressure_{unit}", "temperature_c", "humidity"])
        for ts, p, t, h in rows:
            p = convert_pressure(p, unit)
            w.writerow([ts, fmt_ts(ts, tz, "%Y-%m-%d"), fmt_ts(ts, tz, "%H:%M:%S"),
                        _num(p, 3 if unit == "kPa" else 2, excel_ru), _num(t, 1, excel_ru), _num(h, 1, excel_ru)])
    return len(rows)


def export_json(rows, path, tz, unit="mmHg", meta=None) -> int:
    _check_rows(rows)
    offset = int(to_dt(rows[0][0], tz).utcoffset().total_seconds())
    data = []
    for ts, p, t, h in rows:
        data.append({
            "timestamp": ts,
            "datetime": to_dt(ts, tz).isoformat(timespec="seconds"),
            "pressure": convert_pressure(p, unit),
            "temperature": t,
            "humidity": h,
        })
    doc = {"source": "esp32-barograph", "pressure_unit": unit, "timezone_offset_s": offset,
           "count": len(data), "data": data}
    if meta:
        doc["meta"] = meta
    with open(path, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=1)
    return len(data)


def export_xlsx(rows, path, tz, unit="mmHg", stats=None, title="Барограф ESP32") -> int:
    """Excel: лист «Данные» (настоящие даты/числа), «Статистика» и «Графики» (линейные диаграммы)."""
    _check_rows(rows)
    if len(rows) > MAX_EXCEL_ROWS:
        raise ExportError(f"Слишком много строк для Excel ({len(rows)}). Выберите усреднение (например, 1 час).")
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.axis import DateAxis
    from openpyxl.styles import Alignment, Font

    plabel = pressure_label(unit)
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    headers = ["Время", "Unix время", f"Давление, {plabel}", "Температура, °C", "Влажность, %"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
    has_hum = False
    for ts, p, t, h in rows:
        naive = to_dt(ts, tz).replace(tzinfo=None)         # Excel не хранит часовые пояса
        ws.append([naive, ts, convert_pressure(p, unit), t, h])
        ws.cell(row=ws.max_row, column=1).number_format = "DD.MM.YYYY HH:MM:SS"
        if h is not None:
            has_hum = True
    for col, w in zip("ABCDE", (21, 14, 18, 17, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    n = len(rows)

    # --- статистика
    st = wb.create_sheet("Статистика")
    st.append([title])
    st["A1"].font = Font(bold=True, size=14)
    st.append(["Период", f"{fmt_ts(rows[0][0], tz)} — {fmt_ts(rows[-1][0], tz)}"])
    st.append(["Записей", n])
    st.append([])
    st.append(["Параметр", "Мин", "Макс", "Среднее"])
    for c in st[5]:
        c.font = Font(bold=True)

    def col_stats(idx, conv=None):
        vals = [r[idx] for r in rows if r[idx] is not None]
        if conv:
            vals = [conv(v) for v in vals]
        if not vals:
            return None
        return min(vals), max(vals), sum(vals) / len(vals)

    for name, idx, conv in ((f"Давление, {plabel}", 1, lambda v: convert_pressure(v, unit)),
                            ("Температура, °C", 2, None), ("Влажность, %", 3, None)):
        s = col_stats(idx, conv)
        if s:
            st.append([name] + [round(x, 2) for x in s])
    st.column_dimensions["A"].width = 24
    st.column_dimensions["B"].width = 36

    # --- диаграммы
    cs = wb.create_sheet("Графики")
    specs = [(3, f"Давление, {plabel}"), (4, "Температура, °C")]
    if has_hum:
        specs.append((5, "Влажность, %"))
    for i, (col, ttl) in enumerate(specs):
        ch = LineChart()
        ch.title = ttl
        ch.height, ch.width = 9, 28
        ch.legend = None
        ch.add_data(Reference(ws, min_col=col, min_row=1, max_row=n + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
        ch.y_axis.crossAx = 500
        ch.x_axis = DateAxis(crossAx=100)
        ch.x_axis.number_format = "dd.mm HH:MM"
        ch.x_axis.majorTimeUnit = "days"
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.y_axis.title = ttl
        for s in ch.series:
            s.smooth = False
            s.marker.symbol = "none"
            s.graphicalProperties.line.width = 15000
        cs.add_chart(ch, f"A{1 + i * 20}")
    wb.save(path)
    return n
